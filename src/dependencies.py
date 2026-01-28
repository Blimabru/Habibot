"""Módulo de gerenciamento de dependências do Habibot.

Importa dependências opcionais e registra as que estão faltando.
"""

_DEPENDENCIAS_FALTANDO: list[str] = []

# Selenium imports
try:
    from selenium import webdriver
    from selenium.webdriver.common.by import By
    from selenium.webdriver.common.keys import Keys
    from selenium.webdriver.support.ui import WebDriverWait
    from selenium.webdriver.support import expected_conditions as EC
    from selenium.webdriver.chrome.service import Service
    from selenium.webdriver.chrome.options import Options
    from selenium.common.exceptions import WebDriverException, SessionNotCreatedException, TimeoutException
    from selenium.webdriver import ActionChains
except ModuleNotFoundError:
    webdriver = None
    By = None
    Keys = None
    WebDriverWait = None
    EC = None
    Service = None
    Options = None
    ActionChains = None
    WebDriverException = type('WebDriverException', (Exception,), {})
    SessionNotCreatedException = type('SessionNotCreatedException', (Exception,), {})
    TimeoutException = type('TimeoutException', (Exception,), {})
    _DEPENDENCIAS_FALTANDO.append('selenium')

# WebDriver Manager imports
try:
    from webdriver_manager.chrome import ChromeDriverManager
except ModuleNotFoundError:
    ChromeDriverManager = None
    _DEPENDENCIAS_FALTANDO.append('webdriver_manager')

# OpenPyXL imports
try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, PatternFill, Font, Border, Side
    from openpyxl.drawing.image import Image as OpenpyxlImage
    from openpyxl.utils import get_column_letter
except ModuleNotFoundError:
    Workbook = None
    load_workbook = None
    Alignment = None
    PatternFill = None
    Font = None
    Border = None
    Side = None
    OpenpyxlImage = None
    get_column_letter = None
    _DEPENDENCIAS_FALTANDO.append('openpyxl')


def get_missing_dependencies() -> list[str]:
    """Retorna lista de dependências que estão faltando."""
    return _DEPENDENCIAS_FALTANDO.copy()
